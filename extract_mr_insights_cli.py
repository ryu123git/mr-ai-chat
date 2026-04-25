"""
MR活動報告 → 営業ナレッジ抽出ツール（Claude CLI版）

Claude Codeがインストール済みであればAPIキー不要で動作します。

【実行例】
  # 同フォルダのExcelを処理（デフォルト設定）
  python extract_mr_insights_cli.py --input 所長用の活動報告.xlsx

  # フルパス指定
  python extract_mr_insights_cli.py \
    --input "C:/Users/satou/OneDrive/デスクトップ/所長用の活動報告.xlsx" \
    --output "C:/Users/satou/OneDrive/デスクトップ/活動報告_分析済み.xlsx"

  # 部分処理（100件目から200件目）
  python extract_mr_insights_cli.py --input 所長用の活動報告.xlsx --start 100 --end 200

【オプション】
  --input    入力Excelファイルパス（デフォルト: 所長用の活動報告.xlsx）
  --output   出力Excelファイルパス（デフォルト: 活動報告_分析済み.xlsx）
  --start    処理開始行番号（0始まり）
  --end      処理終了行番号（0始まり・含まない）
  --workers  並列処理数（デフォルト: 3、増やすと高速だが負荷も増加）
  --model    使用するClaudeモデル（デフォルト: claude-haiku-4-5-20251001）
"""

import argparse
import concurrent.futures
import json
import subprocess
import sys
import threading
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from tqdm import tqdm

# ─── 設定 ────────────────────────────────────────────────
ACTIVITY_COL = "活動内容"

NEW_COLUMNS = [
    "他の営業も使える話のネタ",
    "刺さる冒頭トーク",
    "反論に対する切り返し例",
    "反応が鈍い医師には何を確認すべき",
    "エビデンスを求める医師への次の一手",
    "面談シーン",
    "検索キーワード",
    "活動内容サマリー",
]

EXTRACTION_PROMPT_TEMPLATE = """以下のMR活動報告を読んで、各項目に該当する内容があれば抽出してください。
該当しない場合は必ず空文字列を返してください。推測や補完はしないでください。

【活動内容】
{activity}

以下のJSON形式のみで回答してください。説明文・前置き・コードブロックは不要です。
{{
  "他の営業も使える話のネタ": "",
  "刺さる冒頭トーク": "",
  "反論に対する切り返し例": "",
  "反応が鈍い医師には何を確認すべき": "",
  "エビデンスを求める医師への次の一手": "",
  "面談シーン": "",
  "検索キーワード": "",
  "活動内容サマリー": ""
}}

【各項目の抽出基準】
- 他の営業も使える話のネタ: 医師の発言・反応・処方背景など、他MRが話題に使える具体的な情報
- 刺さる冒頭トーク: 医師の興味を引いた切り出し方・アプローチ手法
- 反論に対する切り返し例: 医師の懸念・反論に対してMRが行った効果的な応答
- 反応が鈍い医師には何を確認すべき: 医師が乗り気でない場合に確認すべき障壁・ポイント
- エビデンスを求める医師への次の一手: エビデンス要求に対してMRが提示した・提示すべき対応策
- 面談シーン: 活動内容から読み取れる面談の種類を1つ選択。「初回面談」「継続フォロー」「競合品対応」「エビデンス要求対応」「副作用・安全性対応」「処方拡大依頼」「その他」のいずれか
- 検索キーワード: 活動内容に関連するキーワードをカンマ区切りで列挙（例: 競合品,論文,エビデンス,安全性,副作用）。活動内容から自然に抽出できるものだけ記載
- 活動内容サマリー: 何の話をしたか・医師の反応はどうだったかを1〜2文で簡潔に要約（AIが検索・分類しやすい表現で）"""

# スレッドセーフなカウンター
_lock = threading.Lock()
_error_count = 0


# ─── Claude CLI呼び出し ──────────────────────────────────
def extract_insights(row_idx: int, activity: str, model: str) -> tuple[int, dict]:
    global _error_count
    empty = (row_idx, {col: "" for col in NEW_COLUMNS})

    if not isinstance(activity, str) or not activity.strip():
        return empty

    prompt = EXTRACTION_PROMPT_TEMPLATE.format(activity=activity.strip())

    for attempt in range(3):
        try:
            result = subprocess.run(
                ["claude", "--print", "--model", model, prompt],
                capture_output=True,
                timeout=120,
                encoding="utf-8",
                errors="replace",
            )

            if result.returncode != 0:
                raise RuntimeError(f"claude CLIエラー: {result.stderr[:200]}")

            text = result.stdout.strip()

            # コードブロック除去
            if "```" in text:
                parts = text.split("```")
                for part in parts:
                    part = part.strip()
                    if part.startswith("json"):
                        part = part[4:].strip()
                    if part.startswith("{"):
                        text = part
                        break

            # JSON部分だけ抽出（末尾に説明文があっても対応）
            brace_start = text.find("{")
            brace_end   = text.rfind("}") + 1
            if brace_start >= 0 and brace_end > brace_start:
                text = text[brace_start:brace_end]

            parsed = json.loads(text)
            return (row_idx, {col: str(parsed.get(col, "")).strip() for col in NEW_COLUMNS})

        except json.JSONDecodeError:
            with _lock:
                _error_count += 1
            return empty

        except subprocess.TimeoutExpired:
            if attempt < 2:
                continue
            with _lock:
                _error_count += 1
            return empty

        except Exception as e:
            if attempt < 2:
                continue
            with _lock:
                _error_count += 1
            print(f"\n  [行{row_idx}] エラー: {e}", file=sys.stderr)
            return empty

    return empty


# ─── Excel書き出し ───────────────────────────────────────
def write_output(df: pd.DataFrame, output_path: str):
    df.to_excel(output_path, index=False, sheet_name="活動報告")

    wb = load_workbook(output_path)
    ws = wb.active

    total_cols    = ws.max_column
    new_col_start = total_cols - len(NEW_COLUMNS) + 1

    original_fill = PatternFill("solid", start_color="2E75B6", end_color="2E75B6")
    new_fill      = PatternFill("solid", start_color="375623", end_color="375623")
    orig_font     = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    new_font      = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    body_font     = Font(name="Arial", size=10)
    wrap_top      = Alignment(wrap_text=True, vertical="top")
    center_wrap   = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx in range(1, total_cols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill      = new_fill if col_idx >= new_col_start else original_fill
        cell.font      = new_font if col_idx >= new_col_start else orig_font
        cell.alignment = center_wrap

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font      = body_font
            cell.alignment = wrap_top

    ws.row_dimensions[1].height = 30
    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 70

    special_widths = {"活動日": 14, "活動内容": 50, "ターゲットランク": 16, "使用資材": 12, "売上アップ額": 14}
    headers = [ws.cell(row=1, column=c).value for c in range(1, total_cols + 1)]
    for col_idx, header in enumerate(headers, start=1):
        if col_idx >= new_col_start:
            width = 42
        else:
            width = special_widths.get(header, 18)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    wb.save(output_path)
    print(f"保存完了 → {output_path}")


# ─── メイン ──────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(
        description="MR活動報告から営業ナレッジを抽出してExcelに追記します（Claude CLI版）",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("--input",   default="所長用の活動報告.xlsx",          help="入力Excelファイル")
    parser.add_argument("--output",  default="活動報告_分析済み.xlsx",         help="出力Excelファイル")
    parser.add_argument("--start",   type=int, default=0,                       help="処理開始行（0始まり）")
    parser.add_argument("--end",     type=int, default=None,                    help="処理終了行（0始まり・含まない）")
    parser.add_argument("--workers", type=int, default=3,                       help="並列処理数（デフォルト3）")
    parser.add_argument("--model",   default="claude-haiku-4-5-20251001",       help="使用するClaudeモデル")
    args = parser.parse_args()

    print(f"読み込み中: {args.input}")
    df = pd.read_excel(args.input, sheet_name=0, dtype=str)
    df = df.where(pd.notna(df), "")

    if ACTIVITY_COL not in df.columns:
        sys.exit(
            f"エラー: 列「{ACTIVITY_COL}」が見つかりません。\n"
            f"検出された列: {list(df.columns)}"
        )

    start  = args.start
    end    = args.end if args.end is not None else len(df)
    target = list(range(start, min(end, len(df))))

    print(f"処理対象: {len(target)} 件（全 {len(df)} 件中）")
    print(f"モデル: {args.model}  並列数: {args.workers}")
    print()

    # 新列初期化
    for col in NEW_COLUMNS:
        if col not in df.columns:
            df[col] = ""

    # 並列処理
    results: dict[int, dict] = {}
    with tqdm(total=len(target), desc="抽出中", unit="件") as pbar:
        with concurrent.futures.ThreadPoolExecutor(max_workers=args.workers) as executor:
            futures = {
                executor.submit(extract_insights, idx, df.at[idx, ACTIVITY_COL], args.model): idx
                for idx in target
            }
            for future in concurrent.futures.as_completed(futures):
                row_idx, row_result = future.result()
                results[row_idx] = row_result
                pbar.update(1)

    # データフレームに反映
    for idx in target:
        for col in NEW_COLUMNS:
            df.at[idx, col] = results.get(idx, {}).get(col, "")

    write_output(df, args.output)

    # サマリー
    filled = {col: sum(1 for i in target if df.at[i, col]) for col in NEW_COLUMNS}
    print("\n【抽出件数サマリー】")
    for col, count in filled.items():
        print(f"  {col}: {count} 件 / {len(target)} 件")
    if _error_count:
        print(f"\n  ※ エラー: {_error_count} 件（空白で処理済み）")


if __name__ == "__main__":
    main()
