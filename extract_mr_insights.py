"""
MR活動報告 → 営業ナレッジ抽出ツール

【事前準備】
  pip install anthropic openpyxl pandas tqdm

【APIキーの設定（Windowsコマンドプロンプト）】
  set ANTHROPIC_API_KEY=sk-ant-xxxx

【実行例】
  # 同じフォルダにある活動報告.xlsxを処理
  python extract_mr_insights.py --input 所長用の活動報告.xlsx

  # フルパス指定
  python extract_mr_insights.py --input "C:/Users/satou/OneDrive/デスクトップ/所長用の活動報告.xlsx" --output "C:/Users/satou/OneDrive/デスクトップ/活動報告_分析済み.xlsx"

  # 100件目から200件目だけ処理（再実行・追加処理用）
  python extract_mr_insights.py --input 所長用の活動報告.xlsx --start 100 --end 200

【オプション】
  --input   入力Excelファイルパス（デフォルト: 所長用の活動報告.xlsx）
  --output  出力Excelファイルパス（デフォルト: 活動報告_分析済み.xlsx）
  --start   処理開始行番号（0始まり、デフォルト: 0）
  --end     処理終了行番号（0始まり・含まない、デフォルト: 全件）
  --workers 同時リクエスト数（デフォルト: 5、APIレート制限に合わせて調整）
  --model   使用するClaudeモデル（デフォルト: claude-haiku-4-5-20251001）
"""

import argparse
import asyncio
import json
import os
import sys

import anthropic
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from tqdm import tqdm

# ─── 設定 ────────────────────────────────────────────────
ACTIVITY_COL = "活動内容"

NEW_COLUMNS = [
    "他の営業も使える話のネタ",
    "刺さる冒頭トーク",
    "反論に対する切り返し例",
    "反応が鈍い医師には何を確認すべき",
    "エビデンスを求める医師への次の一手",
    "面談シーン",
    "検索キーワード",
    "活動内容サマリー",
]

SYSTEM_PROMPT = """あなたは製薬会社のMR（医薬情報担当者）の営業コーチです。
MRが記録した活動内容を読み、他のMRが即使えるナレッジを抽出してください。
抽出は事実に基づき、具体的かつ簡潔にまとめてください。"""

EXTRACTION_PROMPT = """以下のMR活動報告を読んで、各項目に該当する内容があれば抽出してください。
該当しない場合は必ず空文字列を返してください。推測や補完はしないでください。

【活動内容】
{activity}

以下のJSON形式のみで回答してください。説明文・前置き・コードブロックは不要です。
{{
  "他の営業も使える話のネタ": "",
  "刺さる冒頭トーク": "",
  "反論に対する切り返し例": "",
  "反応が鈍い医師には何を確認すべき": "",
  "エビデンスを求める医師への次の一手": "",
  "面談シーン": "",
  "検索キーワード": "",
  "活動内容サマリー": ""
}}

【各項目の抽出基準】
- 他の営業も使える話のネタ: 医師の発言・反応・処方背景など、他MRが話題に使える具体的な情報
- 刺さる冒頭トーク: 医師の興味を引いた切り出し方・アプローチ手法
- 反論に対する切り返し例: 医師の懸念・反論に対してMRが行った効果的な応答
- 反応が鈍い医師には何を確認すべき: 医師が乗り気でない場合に確認すべき障壁・ポイント
- エビデンスを求める医師への次の一手: エビデンス要求に対してMRが提示した・提示すべき対応策
- 面談シーン: 活動内容から読み取れる面談の種類を1つ選択。「初回面談」「継続フォロー」「競合品対応」「エビデンス要求対応」「副作用・安全性対応」「処方拡大依頼」「その他」のいずれか
- 検索キーワード: 活動内容に関連するキーワードをカンマ区切りで列挙（例: 競合品,論文,エビデンス,安全性,副作用）。活動内容から自然に抽出できるものだけ記載
- 活動内容サマリー: 何の話をしたか・医師の反応はどうだったかを1〜2文で簡潔に要約（AIが検索・分類しやすい表現で）"""


# ─── Claude API呼び出し ──────────────────────────────────
async def extract_insights(
    client: anthropic.AsyncAnthropic,
    activity: str,
    semaphore: asyncio.Semaphore,
    model: str,
    row_idx: int,
) -> tuple[int, dict]:
    empty = (row_idx, {col: "" for col in NEW_COLUMNS})

    if not isinstance(activity, str) or not activity.strip():
        return empty

    async with semaphore:
        for attempt in range(3):
            try:
                message = await client.messages.create(
                    model=model,
                    max_tokens=1024,
                    system=SYSTEM_PROMPT,
                    messages=[
                        {
                            "role": "user",
                            "content": EXTRACTION_PROMPT.format(activity=activity.strip()),
                        }
                    ],
                )
                text = message.content[0].text.strip()

                # コードブロックがあれば除去
                if "```" in text:
                    parts = text.split("```")
                    for part in parts:
                        part = part.strip()
                        if part.startswith("json"):
                            part = part[4:].strip()
                        if part.startswith("{"):
                            text = part
                            break

                result = json.loads(text)
                return (row_idx, {col: str(result.get(col, "")).strip() for col in NEW_COLUMNS})

            except json.JSONDecodeError as e:
                print(f"\n  [行{row_idx}] JSON解析エラー: {e}", file=sys.stderr)
                return empty

            except anthropic.RateLimitError:
                wait = 60 * (attempt + 1)
                print(f"\n  [行{row_idx}] レート制限 → {wait}秒待機", file=sys.stderr)
                await asyncio.sleep(wait)

            except anthropic.APIError as e:
                print(f"\n  [行{row_idx}] APIエラー: {e}", file=sys.stderr)
                if attempt == 2:
                    return empty
                await asyncio.sleep(5)

    return empty


# ─── バッチ処理（順序保証） ──────────────────────────────
async def run_extraction(activities: list[str], workers: int, model: str) -> list[dict]:
    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        sys.exit("エラー: 環境変数 ANTHROPIC_API_KEY が設定されていません。")

    client = anthropic.AsyncAnthropic(api_key=api_key)
    semaphore = asyncio.Semaphore(workers)

    tasks = [
        extract_insights(client, activity, semaphore, model, i)
        for i, activity in enumerate(activities)
    ]

    # as_completed でプログレスバーを表示しながら収集
    indexed_results: list[tuple[int, dict]] = []
    with tqdm(total=len(tasks), desc="抽出中", unit="件") as pbar:
        for coro in asyncio.as_completed(tasks):
            result = await coro
            indexed_results.append(result)
            pbar.update(1)

    # インデックスで並び替えて順序を復元
    indexed_results.sort(key=lambda x: x[0])
    return [r for _, r in indexed_results]


# ─── Excel書き出し ───────────────────────────────────────
def write_output(df: pd.DataFrame, output_path: str):
    df.to_excel(output_path, index=False, sheet_name="活動報告")

    wb = load_workbook(output_path)
    ws = wb.active

    total_cols = ws.max_column
    new_col_start = total_cols - len(NEW_COLUMNS) + 1

    original_header_fill = PatternFill("solid", start_color="2E75B6", end_color="2E75B6")
    new_header_fill      = PatternFill("solid", start_color="375623", end_color="375623")
    original_header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    new_header_font      = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    body_font            = Font(name="Arial", size=10)
    wrap_top             = Alignment(wrap_text=True, vertical="top")
    center_wrap          = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # ヘッダー行
    for col_idx in range(1, total_cols + 1):
        cell = ws.cell(row=1, column=col_idx)
        if col_idx >= new_col_start:
            cell.fill = new_header_fill
            cell.font = new_header_font
        else:
            cell.fill = original_header_fill
            cell.font = original_header_font
        cell.alignment = center_wrap

    # データ行
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = wrap_top

    # 行高
    ws.row_dimensions[1].height = 30
    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 70

    # 列幅
    col_widths = {
        "活動日":    14,
        "活動内容":  50,
        "ターゲットランク": 16,
        "使用資材":  12,
        "売上アップ額": 14,
    }
    default_width    = 18
    new_column_width = 42

    headers = [ws.cell(row=1, column=c).value for c in range(1, total_cols + 1)]
    for col_idx, header in enumerate(headers, start=1):
        if col_idx >= new_col_start:
            ws.column_dimensions[get_column_letter(col_idx)].width = new_column_width
        else:
            ws.column_dimensions[get_column_letter(col_idx)].width = col_widths.get(header, default_width)

    ws.freeze_panes = "A2"
    wb.save(output_path)
    print(f"保存完了 → {output_path}")


# ─── メイン ──────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(
        description="MR活動報告から営業ナレッジを抽出してExcelに追記します",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("--input",   default="所長用の活動報告.xlsx",          help="入力Excelファイル")
    parser.add_argument("--output",  default="活動報告_分析済み.xlsx",         help="出力Excelファイル")
    parser.add_argument("--start",   type=int, default=0,                       help="処理開始行（0始まり）")
    parser.add_argument("--end",     type=int, default=None,                    help="処理終了行（0始まり・含まない）")
    parser.add_argument("--workers", type=int, default=5,                       help="同時リクエスト数")
    parser.add_argument("--model",   default="claude-haiku-4-5-20251001",       help="使用するClaudeモデル")
    args = parser.parse_args()

    print(f"読み込み中: {args.input}")
    df = pd.read_excel(args.input, sheet_name=0, dtype=str)
    df = df.where(pd.notna(df), "")

    if ACTIVITY_COL not in df.columns:
        sys.exit(
            f"エラー: 列「{ACTIVITY_COL}」が見つかりません。\n"
            f"検出された列: {list(df.columns)}"
        )

    # 処理範囲
    start = args.start
    end   = args.end if args.end is not None else len(df)
    target_indices = list(range(start, min(end, len(df))))

    print(f"処理対象: {len(target_indices)} 件（全 {len(df)} 件中、行 {start}～{end-1}）")
    print(f"モデル: {args.model}  並列数: {args.workers}")

    activities = [df.at[i, ACTIVITY_COL] for i in target_indices]
    results = asyncio.run(run_extraction(activities, args.workers, args.model))

    # 新列を初期化（未処理行は空文字）
    for col in NEW_COLUMNS:
        if col not in df.columns:
            df[col] = ""

    # 結果を書き込み
    for idx, result in zip(target_indices, results):
        for col in NEW_COLUMNS:
            df.at[idx, col] = result.get(col, "")

    write_output(df, args.output)

    # サマリー
    filled_counts = {col: sum(1 for i in target_indices if df.at[i, col]) for col in NEW_COLUMNS}
    print("\n【抽出件数サマリー】")
    for col, count in filled_counts.items():
        print(f"  {col}: {count} 件 / {len(target_indices)} 件")


if __name__ == "__main__":
    main()
